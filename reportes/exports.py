from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from html import escape
from io import BytesIO
from zoneinfo import ZoneInfo

from django.utils import timezone


REPORT_TITLES = {
    'reservas': 'Reservas',
    'cobros': 'Cobros y pagos',
    'cierres': 'Cierres de caja',
    'ocupacion': 'Ocupación de canchas',
}

COLUMNS = {
    'reservas': ('codigo', 'inicio', 'fin', 'cancha', 'deporte', 'cliente', 'estado', 'origen', 'total', 'horas'),
    'cobros': ('numero', 'codigo_pago', 'codigo', 'fecha', 'cancha', 'cliente', 'caja', 'turno', 'operador', 'responsable', 'medio', 'importe', 'estado_reserva'),
    'cierres': ('id', 'caja', 'responsable', 'apertura', 'cierre', 'cerrado_por', 'monto_inicial',
               'cobros_efectivo', 'cobros_electronicos', 'ingresos_manuales', 'retiros', 'esperado', 'contado', 'diferencia'),
    'ocupacion': ('cancha', 'deporte', 'reservas', 'horas_confirmadas', 'horas_pendientes'),
}

NUMERIC = {'id', 'turno', 'total', 'horas', 'importe', 'monto_inicial', 'esperado', 'contado',
           'diferencia', 'reservas', 'horas_confirmadas', 'horas_pendientes', 'cobros_efectivo',
           'cobros_electronicos', 'ingresos_manuales', 'retiros', 'importe_medio'}
INTEGER = {'id', 'turno', 'reservas'}
DATES = {'inicio', 'fin', 'fecha', 'apertura', 'cierre'}
MONEY = {'total', 'importe', 'monto_inicial', 'esperado', 'contado', 'diferencia',
         'cobros_efectivo', 'cobros_electronicos', 'ingresos_manuales', 'retiros', 'importe_medio'}

HEADERS = {
    'id': 'Registro', 'codigo': 'Código de reserva', 'inicio': 'Inicio del servicio', 'fin': 'Fin del servicio',
    'numero': 'Número de recibo', 'codigo_pago': 'Código de pago', 'cancha': 'Cancha', 'deporte': 'Deporte',
    'cliente': 'Cliente', 'estado': 'Estado actual', 'origen': 'Origen', 'total': 'Importe de reserva (Bs)',
    'horas': 'Horas en el período', 'fecha': 'Fecha de cobro', 'caja': 'Caja', 'turno': 'Turno',
    'operador': 'Cobrado por', 'responsable': 'Responsable del turno', 'medio': 'Medios y desglose',
    'importe': 'Total cobrado (Bs)', 'importe_medio': 'Medio seleccionado (Bs)', 'estado_reserva': 'Estado actual de la reserva', 'apertura': 'Apertura del turno',
    'cierre': 'Cierre del turno', 'cerrado_por': 'Cerrado por', 'monto_inicial': 'Fondo inicial (Bs)',
    'esperado': 'Efectivo esperado (Bs)', 'contado': 'Efectivo contado (Bs)', 'diferencia': 'Diferencia (Bs)',
    'reservas': 'Reservas', 'horas_confirmadas': 'Horas confirmadas', 'horas_pendientes': 'Horas pendientes de pago',
    'cobros_efectivo': 'Cobros en efectivo (Bs)', 'cobros_electronicos': 'Cobros electrónicos (Bs)',
    'ingresos_manuales': 'Ingresos manuales (Bs)', 'retiros': 'Retiros (Bs)',
}

SUMMARY_LABELS = {
    'reservas': 'Reservas', 'importe_reservas': 'Importe de reservas (Bs)', 'cobros': 'Cobros',
    'cobrado': 'Total cobrado (Bs)', 'medio_seleccionado': 'Cobrado por medio seleccionado (Bs)', 'efectivo': 'Cobrado en efectivo (Bs)', 'turnos': 'Turnos cerrados',
    'con_diferencia': 'Cierres con diferencia', 'esperado': 'Efectivo esperado (Bs)',
    'contado': 'Efectivo contado (Bs)', 'diferencia': 'Diferencia neta (Bs)',
    'horas_confirmadas': 'Horas confirmadas', 'horas_pendientes': 'Horas pendientes de pago',
    'por_estado': 'Importe de reservas por estado', 'por_medio': 'Cobros por medio',
}

STATE_LABELS = {
    'PENDIENTE': 'Pendiente de pago', 'CONFIRMADA': 'Confirmada', 'EN_CURSO': 'En curso',
    'FINALIZADA': 'Finalizada', 'NO_ASISTIO': 'Cliente no asistió', 'CANCELADA': 'Cancelada', 'EXPIRADA': 'Plazo de pago vencido',
}

ORIGIN_LABELS = {
    'RECEPCION': 'Recepción',
    'INMEDIATO': 'Alquiler inmediato',
    'WEB': 'Solicitud web',
}

NAVY = '102746'
GOLD = 'F5B400'
PALE_BLUE = 'F3F6FA'
TEXT = '203247'
MUTED = '596B80'
WHITE = 'FFFFFF'
GRID = 'D9E2EC'


def columns_for(kind, meta):
    columns = list(COLUMNS[kind])
    if kind == 'cobros' and meta.get('filtros', {}).get('medio'):
        columns.insert(columns.index('importe') + 1, 'importe_medio')
    return columns


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')):
        return "'" + value
    return value


def _summary_pairs(summary):
    pairs = []
    for key, value in summary.items():
        label = SUMMARY_LABELS.get(key, key.replace('_', ' ').title())
        if isinstance(value, dict):
            for child, amount in value.items():
                child_label = STATE_LABELS.get(child, child)
                suffix = ' (Bs)' if key in {'por_estado', 'por_medio'} else ''
                pairs.append((f'{label} - {child_label}{suffix}', amount, False))
        else:
            pairs.append((label, value, key in {'reservas', 'cobros', 'turnos', 'con_diferencia'}))
    return pairs


def _decimal(value):
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return None


def _excel_datetime(value, zone):
    if not value:
        return None
    try:
        result = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
        if timezone.is_naive(result):
            result = result.replace(tzinfo=zone)
        return result.astimezone(zone).replace(tzinfo=None)
    except (ValueError, TypeError):
        return safe_cell(str(value))


def _local_date(value):
    try:
        return date.fromisoformat(str(value)).strftime('%d/%m/%Y')
    except (ValueError, TypeError):
        return str(value)


def display_value(key, value):
    if key in {'estado', 'estado_reserva'}:
        return STATE_LABELS.get(str(value), value)
    if key == 'origen':
        return ORIGIN_LABELS.get(str(value), value)
    return value


def _localized_number(value, decimals=2):
    number = _decimal(value)
    if number is None:
        return str(value)
    formatted = f'{number:,.{decimals}f}'
    return formatted.replace(',', '\0').replace('.', ',').replace('\0', '.')


def pdf_value(key, value, zone):
    value = display_value(key, value)
    if key in DATES:
        converted = _excel_datetime(value, zone)
        if isinstance(converted, datetime):
            return converted.strftime('%d/%m/%Y %H:%M')
        return converted
    if key in NUMERIC:
        return _localized_number(value, 0 if key in INTEGER else 2)
    return value


def make_xlsx(kind, rows, summary, meta):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.properties.title = REPORT_TITLES[kind]
    workbook.properties.subject = f"{meta['sucursal_nombre']} | {meta['desde']} - {meta['hasta']}"
    workbook.properties.creator = 'Soccer Real Madrid'
    summary_sheet = workbook.active
    summary_sheet.title = 'Resumen'
    detail_sheet = workbook.create_sheet('Detalle')
    zone = ZoneInfo(meta['zona_horaria'])

    navy_fill = PatternFill('solid', fgColor=NAVY)
    pale_fill = PatternFill('solid', fgColor=PALE_BLUE)
    white_font = Font(name='Aptos', color=WHITE, bold=True)
    title_font = Font(name='Aptos Display', size=19, bold=True, color=NAVY)
    body_font = Font(name='Aptos', size=10, color=TEXT)
    muted_font = Font(name='Aptos', size=9, color=MUTED)
    thin = Side(style='thin', color=GRID)
    border = Border(bottom=thin)

    summary_sheet.sheet_view.showGridLines = False
    summary_sheet.column_dimensions['A'].width = 38
    summary_sheet.column_dimensions['B'].width = 62
    summary_sheet.merge_cells('A1:B1')
    summary_sheet['A1'] = 'SOCCER REAL MADRID'
    summary_sheet['A1'].font = Font(name='Aptos', size=10, bold=True, color=NAVY)
    summary_sheet['A2'] = REPORT_TITLES[kind]
    summary_sheet['A2'].font = title_font
    summary_sheet['A2'].border = Border(bottom=Side(style='medium', color=GOLD))
    summary_sheet.merge_cells('A2:B2')
    meta_rows = [
        ('Sucursal', meta['sucursal_nombre']),
        ('Período', f"{_local_date(meta['desde'])} al {_local_date(meta['hasta'])}"),
        ('Alcance', 'Sucursal completa' if meta['alcance'] == 'sucursal' else 'Operaciones propias'),
        ('Generado', meta.get('generado_local', meta['generado_en'])),
        ('Generado por', meta.get('generado_por', '')),
        ('Filtros aplicados', meta.get('filtros_texto') or 'Ninguno'),
        ('Criterio del reporte', meta.get('criterio', '')),
    ]
    row_number = 4
    for label, value in meta_rows:
        summary_sheet.cell(row_number, 1, label).font = Font(name='Aptos', size=9, bold=True, color=MUTED)
        summary_sheet.cell(row_number, 2, safe_cell(str(value))).font = body_font
        summary_sheet.cell(row_number, 2).alignment = Alignment(wrap_text=True, vertical='top')
        row_number += 1

    row_number += 1
    for index, label in enumerate(('Indicador', 'Valor'), start=1):
        cell = summary_sheet.cell(row_number, index, label)
        cell.fill = navy_fill
        cell.font = white_font
        cell.alignment = Alignment(vertical='center')
    row_number += 1
    for label, value, is_count in _summary_pairs(summary):
        summary_sheet.cell(row_number, 1, safe_cell(label)).font = body_font
        value_cell = summary_sheet.cell(row_number, 2)
        numeric = _decimal(value)
        if numeric is not None:
            value_cell.value = int(numeric) if is_count and numeric == numeric.to_integral_value() else numeric
            value_cell.number_format = '#,##0' if is_count else '#,##0.00 "Bs"' if '(Bs)' in label else '#,##0.00'
        else:
            value_cell.value = safe_cell(str(value))
        value_cell.font = body_font
        value_cell.alignment = Alignment(horizontal='right' if numeric is not None else 'left')
        for cell in summary_sheet[row_number]:
            cell.border = border
            if row_number % 2 == 0:
                cell.fill = pale_fill
        row_number += 1
    summary_sheet.freeze_panes = 'A13'
    summary_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    summary_sheet.page_setup.orientation = 'portrait'
    summary_sheet.page_setup.fitToWidth = 1
    summary_sheet.page_setup.fitToHeight = 0

    columns = columns_for(kind, meta)
    last_column = len(columns)
    last_letter = get_column_letter(last_column)
    detail_sheet.sheet_view.showGridLines = False
    detail_sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    detail_sheet.cell(1, 1, 'SOCCER REAL MADRID').font = Font(name='Aptos', size=10, bold=True, color=NAVY)
    detail_sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    detail_sheet.cell(2, 1, REPORT_TITLES[kind]).font = title_font
    detail_sheet.cell(2, 1).border = Border(bottom=Side(style='medium', color=GOLD))
    detail_sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_column)
    detail_sheet.cell(3, 1, f"Sucursal: {meta['sucursal_nombre']}    Período: {_local_date(meta['desde'])} al {_local_date(meta['hasta'])}").font = muted_font
    detail_sheet.merge_cells(start_row=4, start_column=1, end_row=4, end_column=last_column)
    detail_sheet.cell(4, 1, f"Generado: {meta.get('generado_local', meta['generado_en'])}    Registros: {len(rows)}").font = muted_font
    header_row = 6
    for column_index, key in enumerate(columns, start=1):
        cell = detail_sheet.cell(header_row, column_index, HEADERS[key])
        cell.fill = navy_fill
        cell.font = white_font
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        cell.border = Border(bottom=Side(style='medium', color=GOLD))
    detail_sheet.row_dimensions[header_row].height = 34

    max_lengths = [len(HEADERS[key]) for key in columns]
    for row_index, item in enumerate(rows, start=header_row + 1):
        for column_index, key in enumerate(columns, start=1):
            raw = item.get(key, '')
            raw = display_value(key, raw)
            cell = detail_sheet.cell(row_index, column_index)
            if raw is None or raw == '':
                cell.value = None
            elif key in DATES:
                cell.value = _excel_datetime(raw, zone)
                if isinstance(cell.value, datetime):
                    cell.number_format = 'dd/mm/yyyy hh:mm'
            elif key in NUMERIC:
                number = _decimal(raw)
                cell.value = int(number) if key in INTEGER and number is not None and number == number.to_integral_value() else number
                if cell.value is not None:
                    cell.number_format = '#,##0.00 "Bs"' if key in MONEY else '#,##0.00' if key not in INTEGER else '#,##0'
            else:
                cell.value = safe_cell(str(raw))
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(vertical='top', wrap_text=key in {'cliente', 'medio', 'cancha', 'deporte', 'responsable', 'operador'})
            if row_index % 2 == 0:
                cell.fill = pale_fill
            text = str(raw) if raw is not None else ''
            max_lengths[column_index - 1] = min(46, max(max_lengths[column_index - 1], min(len(text), 46)))

    for index, key in enumerate(columns, start=1):
        width = max(12, min(max_lengths[index - 1] + 2, 46))
        if key in MONEY or key in NUMERIC:
            width = max(width, 16)
        detail_sheet.column_dimensions[get_column_letter(index)].width = width
    last_row = max(header_row, header_row + len(rows))
    detail_sheet.freeze_panes = f'A{header_row + 1}'
    detail_sheet.auto_filter.ref = f'A{header_row}:{last_letter}{last_row}'
    detail_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    detail_sheet.page_setup.orientation = 'landscape' if last_column > 7 else 'portrait'
    detail_sheet.page_setup.fitToWidth = 1
    detail_sheet.page_setup.fitToHeight = 0
    detail_sheet.print_title_rows = f'{header_row}:{header_row}'
    detail_sheet.sheet_properties.pageSetUpPr.autoPageBreaks = False
    detail_sheet.sheet_properties.outlinePr.summaryBelow = True

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def make_pdf(kind, rows, summary, meta):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    output = BytesIO()
    columns = columns_for(kind, meta)
    zone = ZoneInfo(meta['zona_horaria'])
    page_size = landscape(A4) if len(columns) > 7 else A4
    title = REPORT_TITLES[kind]
    doc = SimpleDocTemplate(output, pagesize=page_size, leftMargin=13 * mm, rightMargin=13 * mm,
        topMargin=22 * mm, bottomMargin=17 * mm, title=title, author='Soccer Real Madrid',
        subject=f"{meta['sucursal_nombre']} | {meta['desde']} - {meta['hasta']}")
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('ReportTitle', parent=styles['Title'], fontName='Helvetica-Bold',
        fontSize=19, leading=22, textColor=colors.HexColor('#102746'), alignment=TA_LEFT, spaceAfter=3 * mm)
    subtitle_style = ParagraphStyle('ReportSubtitle', parent=styles['Normal'], fontName='Helvetica',
        fontSize=9, leading=12, textColor=colors.HexColor('#596B80'), spaceAfter=5 * mm)
    section_style = ParagraphStyle('ReportSection', parent=styles['Heading2'], fontName='Helvetica-Bold',
        fontSize=10, leading=13, textColor=colors.HexColor('#102746'), spaceBefore=3 * mm, spaceAfter=2 * mm)
    body_style = ParagraphStyle('ReportBody', parent=styles['Normal'], fontName='Helvetica',
        fontSize=7.2 if len(columns) > 10 else 8, leading=9, textColor=colors.HexColor('#203247'))
    small_style = ParagraphStyle('ReportSmall', parent=body_style, fontSize=7, leading=8.5)
    metadata_style = ParagraphStyle('ReportMetadata', parent=body_style, fontSize=8, leading=10)
    summary_header_style = ParagraphStyle('ReportSummaryHeader', parent=metadata_style,
        fontName='Helvetica-Bold', textColor=colors.white)
    header_style = ParagraphStyle('ReportHeaderCell', parent=body_style, fontName='Helvetica-Bold',
        fontSize=6.8 if len(columns) > 10 else 7.5, leading=8, textColor=colors.white)

    def para(value, style=body_style, markup=False):
        text = '' if value is None else str(value)
        if not markup:
            text = escape(text).replace('\n', '<br/>')
        return Paragraph(text or '&nbsp;', style)

    def page_chrome(canvas, page_doc):
        canvas.saveState()
        width, height = page_doc.pagesize
        canvas.setFillColor(colors.HexColor('#102746'))
        canvas.rect(0, height - 15 * mm, width, 15 * mm, fill=1, stroke=0)
        canvas.setFillColor(colors.white)
        canvas.setFont('Helvetica-Bold', 9)
        canvas.drawString(page_doc.leftMargin, height - 9.5 * mm, 'SOCCER REAL MADRID')
        canvas.setFont('Helvetica', 7)
        canvas.drawRightString(width - page_doc.rightMargin, height - 9.5 * mm, 'REPORTE OPERATIVO')
        canvas.setStrokeColor(colors.HexColor('#F5B400'))
        canvas.setLineWidth(1.4)
        canvas.line(page_doc.leftMargin, 12 * mm, width - page_doc.rightMargin, 12 * mm)
        canvas.setFillColor(colors.HexColor('#596B80'))
        canvas.setFont('Helvetica', 7)
        canvas.drawString(page_doc.leftMargin, 7.5 * mm, 'Documento generado por el sistema · Uso interno')
        canvas.drawRightString(width - page_doc.rightMargin, 7.5 * mm, f'Página {canvas.getPageNumber()}')
        canvas.restoreState()

    story = [para(title, title_style), para(
        f"Sucursal: {meta['sucursal_nombre']}    Período: {_local_date(meta['desde'])} al {_local_date(meta['hasta'])}", subtitle_style)]
    metadata = [
        [para('<b>Alcance</b>', metadata_style, markup=True), para('Sucursal completa' if meta['alcance'] == 'sucursal' else 'Operaciones propias', metadata_style),
         para('<b>Generado</b>', metadata_style, markup=True), para(meta.get('generado_local', meta['generado_en']), metadata_style)],
        [para('<b>Generado por</b>', metadata_style, markup=True), para(meta.get('generado_por', ''), metadata_style),
         para('<b>Filtros</b>', metadata_style, markup=True), para(meta.get('filtros_texto') or 'Ninguno', metadata_style)],
        [para('<b>Criterio</b>', metadata_style, markup=True), para(meta.get('criterio', ''), metadata_style), '', ''],
    ]
    label_width = 24 * mm
    value_width = (doc.width - 2 * label_width) / 2
    metadata_table = Table(metadata, colWidths=[label_width, value_width, label_width, value_width])
    metadata_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F3F6FA')),
        ('BOX', (0, 0), (-1, -1), .5, colors.HexColor('#D9E2EC')),
        ('INNERGRID', (0, 0), (-1, -1), .35, colors.HexColor('#D9E2EC')),
        ('SPAN', (1, 2), (3, 2)),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6), ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.extend([metadata_table, Spacer(1, 4 * mm), para('Resumen ejecutivo', section_style)])

    summary_rows = [[para('Indicador', summary_header_style), para('Resultado', summary_header_style)]]
    for label, value, is_count in _summary_pairs(summary):
        if _decimal(value) is not None:
            display = _localized_number(value, 0 if is_count else 2)
            display_value = f'Bs {display}' if not is_count and '(Bs)' in label else display
        else:
            display_value = str(value)
        summary_rows.append([para(label, metadata_style), para(display_value, metadata_style)])
    summary_table = Table(summary_rows, colWidths=[doc.width * .66, doc.width * .34], repeatRows=1)
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#102746')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F6FA')]),
        ('BOX', (0, 0), (-1, -1), .5, colors.HexColor('#D9E2EC')),
        ('INNERGRID', (0, 0), (-1, -1), .35, colors.HexColor('#D9E2EC')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6), ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.extend([summary_table, Spacer(1, 3 * mm), para(f'Detalle del reporte · {len(rows)} registros', section_style)])

    table_data = [[para(HEADERS[key], header_style) for key in columns]]
    for row in rows:
        table_data.append([para(pdf_value(key, row.get(key, ''), zone), small_style) for key in columns])
    weights = [1.5 if key in {'cliente', 'medio'} else 1.2 if key in {'cancha', 'responsable', 'operador'} else .78 if key in NUMERIC else 1 for key in columns]
    total_weight = sum(weights)
    col_widths = [doc.width * weight / total_weight for weight in weights]
    detail_table = LongTable(table_data, colWidths=col_widths, repeatRows=1, splitByRow=1, hAlign='LEFT')
    detail_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#102746')),
        ('LINEBELOW', (0, 0), (-1, 0), 1.4, colors.HexColor('#F5B400')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F6FA')]),
        ('GRID', (0, 0), (-1, -1), .3, colors.HexColor('#D9E2EC')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4), ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(detail_table)
    doc.build(story, onFirstPage=page_chrome, onLaterPages=page_chrome)
    return output.getvalue()
